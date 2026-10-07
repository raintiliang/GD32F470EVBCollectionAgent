import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_report():
    doc = docx.Document()
    
    # Title
    title = doc.add_heading('AEX Manager 软件评估与测试进度报告', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 1. Background
    doc.add_heading('一、 项目背景与软件现状', level=1)
    doc.add_paragraph('目前已完成 JTAG 测试基础环境的搭建：')
    doc.add_paragraph('软件安装：已安装 AEX Manager 软件并完成初步功能验证。', style='List Bullet')
    doc.add_paragraph('硬件连接：基于 JT37x7 Blaster 硬件进行物理连接与简单电路测试。', style='List Bullet')
    doc.add_paragraph('流程掌握：已熟悉 AEX Manager 界面操作及基础项目创建与测试流程。', style='List Bullet')

    # 2. Project Files
    doc.add_heading('二、 DB_HRN4_REV1 项目开发进度', level=1)
    doc.add_paragraph('针对张总提供的 DB_HRN4_REV1 项目，目前处于设计文件适配阶段：')
    doc.add_paragraph('已接收资料：Netlist 压缩包、BOM 文件、DB_HRN4_REV1.gen 文件、tk1_top.bsdl、DB_HRN4_Rev1_orPdump64.net, DB_HRN4_Rev1_orPdump64_provision.net 等。', style='List Bullet')
    doc.add_paragraph('适配结论：经核实，上述网表与网表/BOM 文件无法直接导入 AEX Manager 软件上进行测试。', style='List Bullet')

    # 3. Vendor Discussion
    doc.add_heading('三、 JTAG Technologies (中国) 技术沟通反馈', level=1)
    doc.add_paragraph('经与 JTAG 中国公司技术专家沟通讨论，明确了以下关键点：')
    doc.add_paragraph('工具链依赖：AEX Manager 为执行层软件，电路板的测试序列开发需要安装使用 ProVision 软件。', style='List Bullet')
    doc.add_paragraph('BSDL 文件冲突：技术人员审查发现 tk1_top.bsdl 文件与电路图中 SoC 的引脚封装不一致。', style='List Bullet')

    # 4. Alternative Verification
    doc.add_heading('四、 交叉验证与问题排查', level=1)
    doc.add_paragraph('团队使用手头其他具备 JTAG 电路的板卡进行了替代方案验证：')
    doc.add_paragraph('测试工具：采用 JTAG Live 免费版本（类似 ProVision 的低配版软件）尝试开发 AEX Manager 测试序列。', style='List Bullet')
    doc.add_paragraph('测试结果：开发出来的项目存在问题。使用 AEX Manager 对该电路板进行测试时，基础电路测试（Infrastructure Test）失败。', style='List Bullet')

    # 5. Conclusions and Recommendations
    doc.add_heading('五、 结论与后续工作计划', level=1)
    doc.add_paragraph('针对以上问题，后续建议按以下方向推进：')
    doc.add_paragraph('1. 软件环境：联系厂商获取 ProVision 软件授权，用于开发 DB_HRN4_REV1 的完整测试序列。', style='List Number')
    doc.add_paragraph('2. 文件修正：联系芯片厂商获取与电路图 SoC 封装完全匹配的准确 BSDL 文件。', style='List Number')
    doc.add_paragraph('3. 序列调试：配合规范的 ProVision 环境和正确 BSDL，重新开发测试项目，解决基础电路测试失败的问题。', style='List Number')

    doc.save('AEX_Manager_Report_霆宇.docx')
    print("Report generated successfully.")

if __name__ == "__main__":
    create_report()
