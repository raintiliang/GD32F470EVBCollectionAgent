import os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_proposal():
    doc = Document()
    
    # Title
    title = doc.add_heading('材料测试设备及软件系统方案建议书', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Project Overview
    doc.add_heading('1. 项目概述', level=1)
    doc.add_paragraph(
        "本项目旨在开发一套集成了环境监控、自动控制、本地存储及云端远程管理的材料测试设备。"
        "该设备需在特定的二氧化碳（CO2）浓度和湿度环境下维持材料的持续反应，并支持多端访问和历史数据查询。"
    )
    
    # Hardware Suggestion
    doc.add_heading('2. 硬件方案建议 (Hardware)', level=1)
    
    p = doc.add_paragraph()
    p.add_run('核心主控：建议选用 GD32F470 系列').bold = True
    doc.add_paragraph(
        "理由：GD32F470 拥有更高主频 (240MHz)、专用的 LTDC 显示控制器和硬件图形加速（IPA），"
        "能够流畅驱动 5 寸 RGB 屏幕 (800x480) 并处理复杂的数据通讯。相比 GD32G553，更适合作为网关级主控。"
    )
    
    # Hardware Table
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = '功能模块'
    hdr_cells[1].text = '接口类型'
    hdr_cells[2].text = '建议选型/说明'
    
    items = [
        ['二氧化碳传感器', 'UART', 'NDIR 原理传感器 (如 MH-Z19C)'],
        ['温湿度传感器', 'I2C', '工业级传感器 (如 SHT30/SHT40)'],
        ['电磁阀控制', 'GPIO', '需配继电器或驱动电路 (控制CO2输送)'],
        ['除湿机模组', 'RS485', '需加 TTL 转 RS485 转换芯片'],
        ['加湿机模块', 'GPIO', '需配继电器驱动'],
        ['数据存储', 'SDIO', 'TF 卡 (支持 FATFS 文件系统)'],
        ['无线通讯', 'UART', '4G Cat.1 模块 (如 Air780E/EC200S)'],
        ['本地显示', 'RGB/LTDC', '5 寸 TFT LCD (800x480) + LVGL GUI'],
        ['用户交互', 'GPIO', '机械按键或电容触摸屏']
    ]
    
    for mod, iface, desc in items:
        row_cells = table.add_row().cells
        row_cells[0].text = mod
        row_cells[1].text = iface
        row_cells[2].text = desc

    # Software Architecture
    doc.add_heading('3. 软件架构方案 (Software)', level=1)
    
    doc.add_heading('3.1 设备端 (MCU Firmware)', level=2)
    doc.add_paragraph("• 操作系统：RT-Thread 或 FreeRTOS。", style='List Bullet')
    doc.add_paragraph("• UI 框架：LVGL (用于实现实时数值显示及 TF 卡内曲线绘制)。", style='List Bullet')
    doc.add_paragraph("• 存储逻辑：数据本地存入 TF 卡 (CSV/二进制)，定时同步至云端。", style='List Bullet')
    doc.add_paragraph("• 通讯协议：MQTT (经 4G 模块与服务器通讯)。", style='List Bullet')
    
    doc.add_heading('3.2 云端与网页 (Cloud & Web)', level=2)
    doc.add_paragraph("• 后端：Node.js 或 Python (FastAPI) 处理逻辑。")
    doc.add_paragraph("• 数据库：InfluxDB (存储时序曲线) + MySQL (存储用户信息)。")
    doc.add_paragraph("• 前端：Vue.js / React 响应式网页，适配 PC 和手机。")
    doc.add_paragraph("• 图表：ECharts 实现历史数据的动态曲线展示。")
    
    # Cost Estimation
    doc.add_heading('4. 费用与周期估算 (Cost & Timeline)', level=1)
    
    cost_table = doc.add_table(rows=1, cols=3)
    cost_table.style = 'Table Grid'
    c_hdr = cost_table.rows[0].cells
    c_hdr[0].text = '开发模块'
    c_hdr[1].text = '周期'
    c_hdr[2].text = '预估费用 (RMB)'
    
    costs = [
        ['嵌入式固件开发', '4-6 周', '2.5万 - 4.5万'],
        ['云平台后端开发', '2-3 周', '1.5万 - 2.5万'],
        ['响应式网页前端', '2-3 周', '1.0万 - 2.0万'],
        ['合计', '8-12 周', '5.0万 - 9.0万']
    ]
    
    for m, p, f in costs:
        row_cells = cost_table.add_row().cells
        row_cells[0].text = m
        row_cells[1].text = p
        row_cells[2].text = f

    doc.add_paragraph("\n*注：以上费用为行业参考外包价格，实际成本随功能细节调整。")
    
    file_path = '材料测试设备方案建议书.docx'
    doc.save(file_path)
    print(file_path)

if __name__ == "__main__":
    create_proposal()
