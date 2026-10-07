import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

def create_exhibition_xlsx():
    wb = openpyxl.Workbook()
    
    # Sheet 1: Energy Industry (能源电力)
    ws1 = wb.active
    ws1.title = "能源电力领域"
    
    headers = ["领域", "展会名称", "地点", "时间 (预计)", "官方网址", "主题 focus (工艺过程自动控制)", "重点品牌关注"]
    
    # Data for Energy
    energy_data = [
        ["电力能源", "EP China 2026 (国际电力电工展)", "上海", "2026年12月", "www.epchinashow.com", "智能电网自动化、电厂控制系统、继电保护、能量管理系统(EMS)", "和利时(电力自动化强项)、中控技术"],
        ["电力能源", "GPower 2026 (动力展及发电机组展)", "上海", "2026年6月", "www.gpower-expo.com", "分布式能源控制、发电机组自动化控制、并网技术", "中控技术、和利时"],
        ["核能", "Nuclear Industry China 2026 (国际核工业展)", "北京", "2026年4月 (已结束) / 2028", "www.nic-expo.net", "核电站仪控系统(I&C)、DCS平台、核级数字化控制", "和利时(核电控制领军)、中控技术"],
        ["新能源/储能", "SNEC 2026 (光伏及储能展)", "上海", "2026年6月", "www.snec.org.cn", "储能系统集成、PCS控制、能源互联网平台", "中控技术(S-iPower)"],
        ["电力能源", "EP China 2027", "北京/上海", "2027年11月", "www.epchinashow.com", "源网荷储协同控制、数字化变电站、虚拟电厂控制", "和利时、中控技术"],
    ]
    
    # Sheet 2: Petrochemical Industry (石油化工)
    ws2 = wb.create_sheet("石油化工领域")
    
    petro_data = [
        ["石油化工", "CIPPE 2027 (中国国际石油石化展)", "北京", "2027年3月", "www.cippe.com.cn", "数字化油田、炼化一体化DCS、油气管道SCADA、SIS安全系统", "中控技术(流程工业主力)、和利时"],
        ["石油化工", "CIOAE 2027 (石油化工自动化展)", "北京", "2027年3月", "www.cippe.com.cn", "自动化仪表、智能炼厂、工控安全、流程优化软件", "中控技术、和利时"],
        ["精细化工", "ICIF China 2026 (国际化工展览会)", "上海", "2026年9月", "www.icif.cn", "化工过程自动化、安全生产管理系统、智能工厂解决方案", "中控技术、和利时"],
        ["石油化工", "CIPPE 2026 (已结束)", "北京", "2026年3月", "www.cippe.com.cn", "油气开采自动化、海洋工程控制、防爆仪表", "中控技术、和利时"],
    ]

    # Sheet 3: General Process Automation (自动化/仪表 - 品牌核心)
    ws3 = wb.create_sheet("核心自动化/仪表展")
    ws3_data = [
        ["仪表控制", "MICONEX 2026 (中国国际测量控制与仪器仪表展)", "预计北京/成都", "2026年10-11月", "www.miconex.com.cn", "中控技术与和利时的主场。涵盖DCS、PLC、工业软件、变送器", "中控技术、和利时(核心展示)"],
        ["工业自动化", "IAS 2026 (工业自动化展 - 工博会)", "上海", "2026年9月", "www.industrial-automation-show.com", "流程工业数字化转型、边缘计算、自主PLC/DCS系统", "中控技术、和利时"],
        ["过程工业", "Flowexpo 2027 (广州流量展)", "广州", "2027年5月", "www.flowexpo.org", "阀门定位器、流量控制自动化、过程管理", "中控技术"],
    ]

    def format_sheet(ws, data):
        ws.append(headers)
        for row in data:
            ws.append(row)
        
        # Styling
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
            for cell in row:
                cell.border = thin_border
                cell.alignment = Alignment(wrap_text=True, vertical="center")
        
        # Column widths
        ws.column_dimensions['A'].width = 15
        ws.column_dimensions['B'].width = 40
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 20
        ws.column_dimensions['E'].width = 25
        ws.column_dimensions['F'].width = 50
        ws.column_dimensions['G'].width = 30

    format_sheet(ws1, energy_data)
    format_sheet(ws2, petro_data)
    format_sheet(ws3, ws3_data)
    
    wb.save("China_Energy_Petro_Automation_Exhibitions_2026-2027.xlsx")

if __name__ == "__main__":
    create_exhibition_xlsx()
